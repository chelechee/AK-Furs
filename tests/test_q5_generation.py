import re
import shutil
from copy import deepcopy
from pathlib import Path

import pytest
from docx import Document

from akfurs import q5
from akfurs.io_loaders import load_dataset
from akfurs.normalize import order_no
from akfurs.orders import build_clean_lines, summarise
from akfurs.report import build_answers
from tools import build_q5
from tools.check_q5 import verify


ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope='module')
def dataset(data_dir):
    return load_dataset(data_dir)


@pytest.fixture(scope='module')
def figures(dataset):
    return q5.build_figures(dataset)


def _document_date():
    markdown = (ROOT / 'Q5.md').read_text(encoding='utf-8')
    return re.findall(r'AK Furs · [^·]+ · (\d{2}\.\d{2}\.\d{4})', markdown)[0]


def test_figures_come_from_data(figures):
    dataset = load_dataset(ROOT / 'data')
    q1_orders, q1_lines = summarise(
        build_clean_lines(dataset['pricebook'], dataset['crm']))
    answers = build_answers(dataset)
    assert figures['raw_rows'] == len(dataset['crm'])
    assert (figures['unique_orders'], figures['unique_lines']) == (q1_orders, q1_lines)
    assert figures['duplicate_lines'] == len(dataset['crm']) - q1_lines
    assert figures['paid_without_payment'] == len(answers['q2a_paid_orders_without_payment'])
    assert figures['unidentified_count'] == answers['q2b_unidentified_payments_count']
    assert figures['violations_total'] == answers['q3_price_violation_lines_total']
    assert figures['flagged_chats'] == len(answers['q4_chats_meeting_before_price'])
    assert figures['chats_total'] == len(dataset['chats'])
    known_orders = {line.order_no for line in build_clean_lines(
        dataset['pricebook'], dataset['crm'])}
    assert figures['chat_order_missing_count'] == sum(
        bool(chat['order_raw']) and order_no(chat['order_raw']) not in known_orders
        for chat in dataset['chats'])


def test_zero_findings_are_reported_without_claiming_problems(figures):
    text = '\n'.join(q5.memo_paragraphs(figures, '23.10.2026'))
    if not figures['duplicate_lines']:
        assert 'Дубли строк не обнаружены' in text
        assert 'объём завышается' not in text
    if not figures['paid_without_payment'] and not figures['unidentified_count']:
        assert 'Расхождений по сверке поступлений не выявлено' in text
        assert 'Деньги сходятся не полностью' not in text
    if not figures['violations_total']:
        assert 'Продаж ниже минимально допустимых цен не выявлено' in text


def test_zero_amount_formatting_is_safe():
    assert q5.millions(0) == '0 ₽'
    assert q5.share(0, 0) == '0%'


def test_generated_markdown_matches_data(figures, tmp_path):
    date = '23.10.2026'
    expected = q5.markdown(figures, date)
    assert q5.unexpected_numbers(expected, q5.allowed_numbers(figures)) == []
    assert verify(expected, q5.docx_texts(figures, date), figures) == date


def test_generated_docx_texts_match_rendered_data(figures, tmp_path):
    date = '23.10.2026'
    expected = [text.strip() for text in q5.docx_texts(figures, date)]
    path = tmp_path / 'generated.docx'
    build_q5.write_q5(tmp_path / 'generated.md', path, q5.markdown(figures, date),
                      q5.docx_texts(figures, date))
    actual = [paragraph.text.strip() for paragraph in Document(path).paragraphs]
    assert actual == expected


def test_memo_numbers_follow_data(figures):
    date = _document_date()
    changed = dict(figures)
    changed['paid_without_payment'] = 11
    changed['violations_total'] = 70
    text = '\n'.join(q5.memo_paragraphs(changed, date))
    assert '11 заказов со статусом «Оплачен»' in text
    assert '70 позиций проданы ниже минимально допустимой цены' in text
    assert 'Сделать номер заказа обязательным' in text
    original = '\n'.join(q5.memo_paragraphs(figures, date))
    assert '11 заказов со статусом' not in original


def test_client_reply_numbers_follow_data(figures):
    date = _document_date()
    changed = {'client': dict(figures['client'])}
    changed['client']['crm_link_matches'] = True
    changed['client']['chat_present'] = True
    changed['client']['item_name'] = 'жилет соболь натуральный'
    changed['client']['discount_kind'] = None
    changed['client']['price'] = '999 000'
    changed['client']['total'] = '899 100'
    reply = '\n'.join(q5.reply_paragraphs(changed, date))
    assert '999 000 ₽' in reply
    assert '899 100 ₽' in reply
    assert '252 000 ₽' not in reply


def test_client_reply_respects_chat_crm_link_status(figures):
    client = figures['client']
    reply = '\n'.join(q5.reply_paragraphs(figures, '23.10.2026'))
    assert client['chat_present']
    if client['crm_link_matches']:
        assert client['price'] is not None or client['total'] is not None
        if client['price'] is not None:
            assert client['price'] in reply
        if client['discount_percent'] == '10' and client['total'] == '226 800':
            assert 'Жилет из натурального соболя стоит 252 000 ₽.' in reply
            assert 'С учётом вашей скидки 10% стоимость составит 226 800 ₽.' in reply
            assert 'оценить оттенок меха и фактуру изделия при дневном свете' in reply
            assert 'ежедневно с 10:00 до 21:00' in reply
            assert 'когда вам было бы удобно приехать?' in reply
            assert 'модели ✨' in reply
    else:
        assert client['price'] is None
        assert 'изделия, о котором вы спрашиваете' in reply
        assert '₽' not in reply
        memo = '\n'.join(q5.memo_paragraphs(figures, '23.10.2026'))
        assert 'Связь чата C0010 с записью CRM не подтверждается' in memo
    assert 'Справочно для проверки' not in reply


def test_mismatched_chat_customer_never_gets_crm_price(dataset):
    dataset = dict(dataset)
    lines = build_clean_lines(dataset['pricebook'], dataset['crm'])
    line = lines[0]
    item = next(row['name'] for row in dataset['pricebook']['catalog']
                if row['article'] == line.article)
    chat = deepcopy(dataset['chats'][0])
    chat['chat_id'] = 'C0010'
    chat['order_raw'] = 'A-{:05d}'.format(line.order_no)
    chat['client_raw'] = 'different-client-id'
    chat['messages'] = [
        {**message, 'text': item if message['sender'] == 'Клиент' else message['text']}
        for message in chat['messages']
    ]
    dataset['chats'] = [chat]
    figures = q5.build_figures(dataset)
    assert figures['client']['chat_present']
    assert not figures['client']['crm_link_matches']
    assert figures['client']['price'] is None
    reply = '\n'.join(q5.reply_paragraphs(figures, '23.10.2026'))
    assert figures['client']['item_name'] == item
    assert 'изделия, о котором вы спрашиваете' in reply
    assert '₽' not in reply
    memo = '\n'.join(q5.memo_paragraphs(figures, '23.10.2026'))
    assert 'Связь чата C0010 с записью CRM не подтверждается' in memo


def test_chat_not_present_never_generates_a_client_message(figures):
    if figures['client']['chat_present']:
        pytest.skip('C0010 is present in this dataset')
    text = '\n'.join(q5.reply_paragraphs(figures, '23.10.2026'))
    assert 'Ответ не сформирован' in text
    assert 'Не отправлять клиентке' in text


def test_static_text_has_no_digits():
    for text in (q5.MEMO_THEME, q5.MEMO_GREETING, q5.PROPOSALS_TITLE,
                 q5.CLIENT_GREETING, q5.CLIENT_QUESTION,
                 q5.MEMO_TITLE, q5.REPLY_TITLE):
        assert not re.search(r'\d', re.sub(r'C\d+', '', text)), text
    assert q5.BUSINESS_HOURS == (10, 21, 0)


def test_generated_text_has_no_numbers_outside_data(figures):
    date = _document_date()
    text = q5.markdown(figures, date)
    assert q5.unexpected_numbers(text, q5.allowed_numbers(figures)) == []


def test_stale_number_is_detected(figures):
    date = _document_date()
    count = figures['raw_rows']
    text = re.sub(
        r'(?<!\d){}(?=\s+строк)'.format(count),
        '777', q5.markdown(figures, date), count=1)
    with pytest.raises(ValueError, match='does not match the data'):
        verify(text, q5.docx_texts(figures, date), figures)
    assert '777' in q5.unexpected_numbers(text, q5.allowed_numbers(figures))


def test_missing_second_date_is_rejected(figures):
    date = _document_date()
    text = q5.markdown(figures, date).replace('AK Furs · ответ менеджера · ' + date,
                                              'AK Furs · ответ менеджера · без даты')
    with pytest.raises(ValueError, match='two dated lines'):
        verify(text, q5.docx_texts(figures, date), figures)


def test_docx_update_rewrites_texts_and_keeps_structure(figures, tmp_path):
    target = tmp_path / 'Q5.docx'
    shutil.copyfile(ROOT / 'Q5.docx', target)
    new_date = '01.01.2027'
    expected_texts = q5.docx_texts(figures, new_date)
    build_q5.update_docx(target, expected_texts)
    paragraphs = [paragraph.text for paragraph in Document(target).paragraphs]
    assert len(paragraphs) == len(expected_texts)
    assert paragraphs[1] == 'AK Furs · итоги сверки данных · ' + new_date
    assert paragraphs[13] == 'AK Furs · ответ менеджера · ' + new_date


def test_q5_docx_uses_consistent_business_typography(figures, tmp_path):
    target = tmp_path / 'styled.docx'
    texts = q5.docx_texts(figures, '23.10.2026')
    build_q5.write_q5(tmp_path / 'styled.md', target, q5.markdown(figures, '23.10.2026'), texts)
    document = Document(target)
    assert document.sections[0].page_width.cm == pytest.approx(21.0, abs=0.002)
    assert document.paragraphs[0].runs[0].font.name == 'Arial'
    assert document.paragraphs[0].runs[0].font.size.pt == 15
    assert document.paragraphs[1].runs[0].font.size.pt == 9
    assert document.paragraphs[12].paragraph_format.page_break_before
    assert all(run.font.name == 'Arial' for paragraph in document.paragraphs
               for run in paragraph.runs)


def test_docx_update_rejects_wrong_structure(figures, tmp_path):
    target = tmp_path / 'broken.docx'
    document = Document()
    document.add_paragraph('Записка руководителю')
    document.save(target)
    with pytest.raises(ValueError, match='expected'):
        build_q5.update_docx(target, q5.docx_texts(figures, '01.01.2027'))
