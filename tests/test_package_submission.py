import json
import shutil

import pytest
from docx import Document

from tools import package_submission


ROOT = package_submission.ROOT


def _make_project(tmp_path, monkeypatch, with_letter=True, raw_file=None):
    project = tmp_path / 'project'
    handoff = project / 'handoff'
    package = project / 'employer_package'
    handoff.mkdir(parents=True)
    package.mkdir()

    shutil.copyfile(ROOT / 'answers.json', project / 'answers.json')
    shutil.copyfile(ROOT / 'answers.json', handoff / 'answers.json')
    for name in ('Q5.md', 'Q5.docx'):
        shutil.copyfile(ROOT / name, project / name)
        shutil.copyfile(ROOT / name, handoff / name)

    if with_letter:
        for name in package_submission.LETTER_FILES:
            (package / name).write_text('черновик', encoding='utf-8')
        (package / package_submission.SCREENSHOTS_DIR).mkdir()
        (package / package_submission.SCREENSHOTS_DIR / 'shot.jpg').write_bytes(b'jpeg')
    if raw_file:
        (package / raw_file).write_bytes(b'raw')

    monkeypatch.setattr(package_submission, 'ROOT', project)
    monkeypatch.setattr(package_submission, 'HANDOFF', handoff)
    monkeypatch.setattr(package_submission, 'PACKAGE', package)
    return project, handoff, package


def test_package_copies_answers_and_word_only(tmp_path, monkeypatch):
    project, handoff, package = _make_project(tmp_path, monkeypatch)

    answers = package_submission.build_package()

    packaged = json.loads((package / 'answers.json').read_text(encoding='utf-8'))
    assert packaged == answers
    assert packaged['q4_chats_meeting_before_price'] == json.loads(
        (handoff / 'answers.json').read_text(encoding='utf-8'))['q4_chats_meeting_before_price']
    assert (package / 'Q5.docx').is_file()
    assert not (package / 'Q5.md').exists()

    paragraphs = [paragraph.text for paragraph in Document(package / 'Q5.docx').paragraphs]
    assert sum(text.startswith('AK Furs · ') for text in paragraphs) == 2
    assert any(text.startswith(('Здравствуйте!', 'Ответ не сформирован:'))
               for text in paragraphs)


def test_package_uses_current_handoff_q5_not_the_project_snapshot(tmp_path, monkeypatch):
    _, _, package = _make_project(tmp_path, monkeypatch)
    markdown = package_submission.HANDOFF / 'Q5.md'
    markdown.write_text(
        markdown.read_text(encoding='utf-8').replace(
            'Здравствуйте! Спасибо за ваш интерес к нашей модели ✨',
            'Здравствуйте! Проверка handoff ✨'),
        encoding='utf-8',
    )
    package_submission.build_package()
    texts = [paragraph.text for paragraph in Document(package / 'Q5.docx').paragraphs]
    assert any('Проверка handoff' in text for text in texts)


def test_package_main_warns_when_client_reply_cannot_be_generated(tmp_path, monkeypatch, capsys):
    _make_project(tmp_path, monkeypatch)
    markdown = package_submission.HANDOFF / 'Q5.md'
    markdown.write_text(
        markdown.read_text(encoding='utf-8').replace(
            'Здравствуйте! Спасибо за ваш интерес к нашей модели ✨',
            'Ответ не сформирован: переписка C0010 отсутствует;'),
        encoding='utf-8',
    )
    assert package_submission.main() == 0
    assert 'must not be sent' in capsys.readouterr().out


def test_package_requires_letter_materials(tmp_path, monkeypatch):
    _make_project(tmp_path, monkeypatch, with_letter=False)
    with pytest.raises(FileNotFoundError, match='employer_package is missing'):
        package_submission.build_package()


def test_package_refuses_raw_inputs(tmp_path, monkeypatch):
    _make_project(tmp_path, monkeypatch, raw_file='orders.csv')
    with pytest.raises(ValueError, match='must not contain raw inputs'):
        package_submission.build_package()


def test_package_requires_handoff_answers(tmp_path, monkeypatch):
    project, handoff, package = _make_project(tmp_path, monkeypatch)
    (handoff / 'answers.json').unlink()
    with pytest.raises(FileNotFoundError, match='handoff/answers.json'):
        package_submission.build_package()
