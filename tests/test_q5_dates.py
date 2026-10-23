import shutil
from pathlib import Path

from docx import Document

from tools.refresh_q5_date import PREFIXES, missing_prefixes, update_docx, update_markdown


ROOT = Path(__file__).resolve().parents[1]
TODAY = '23.10.2026'


def test_markdown_refresh_updates_memo_and_reply_dates(tmp_path):
    path = tmp_path / 'Q5.md'
    path.write_text(
        'AK Furs · итоги сверки данных · 01.01.2025\n'
        'AK Furs · ответ менеджера · 02.01.2025\n',
        encoding='utf-8',
    )
    assert update_markdown(path, TODAY) is True
    text = path.read_text(encoding='utf-8')
    assert 'итоги сверки данных · {}'.format(TODAY) in text
    assert 'ответ менеджера · {}'.format(TODAY) in text
    assert missing_prefixes(text) == []
    assert update_markdown(path, TODAY) is False


def test_docx_refresh_updates_memo_and_reply_dates(tmp_path):
    path = tmp_path / 'Q5.docx'
    document = Document()
    document.add_paragraph('AK Furs · итоги сверки данных · 01.01.2025')
    document.add_paragraph('AK Furs · ответ менеджера · 02.01.2025')
    document.save(path)

    assert update_docx(path, TODAY) is True
    paragraphs = [paragraph.text for paragraph in Document(path).paragraphs]
    assert paragraphs == [
        'AK Furs · итоги сверки данных · {}'.format(TODAY),
        'AK Furs · ответ менеджера · {}'.format(TODAY),
    ]
    assert update_docx(path, TODAY) is False


def test_single_command_updates_both_files_and_keeps_them_in_sync(tmp_path):
    markdown = tmp_path / 'Q5.md'
    docx = tmp_path / 'Q5.docx'
    markdown.write_text(
        'ЗАПИСКА РУКОВОДИТЕЛЮ\n\n'
        'AK Furs · итоги сверки данных · 01.01.2025\n\n'
        'ОТВЕТ КЛИЕНТКЕ В ЧАТЕ C0010\n\n'
        'AK Furs · ответ менеджера · 01.01.2025\n\n'
        'Здравствуйте! Текст ответа.\n',
        encoding='utf-8',
    )
    document = Document()
    document.add_paragraph('AK Furs · итоги сверки данных · 01.01.2025')
    document.add_paragraph('AK Furs · ответ менеджера · 01.01.2025')
    document.add_paragraph('Здравствуйте! Текст ответа.')
    document.save(docx)

    update_markdown(markdown, TODAY)
    update_docx(docx, TODAY)

    markdown_dates = [line.strip() for line in markdown.read_text(encoding='utf-8').splitlines()
                      if line.strip().startswith(PREFIXES)]
    docx_dates = [paragraph.text.strip() for paragraph in Document(docx).paragraphs
                  if paragraph.text.strip().startswith(PREFIXES)]
    assert markdown_dates == docx_dates == [
        'AK Furs · итоги сверки данных · {}'.format(TODAY),
        'AK Furs · ответ менеджера · {}'.format(TODAY),
    ]


def test_project_q5_files_have_both_dates(tmp_path):
    for name in ('Q5.md', 'Q5.docx'):
        shutil.copyfile(ROOT / name, tmp_path / name)

    update_markdown(tmp_path / 'Q5.md', TODAY)
    update_docx(tmp_path / 'Q5.docx', TODAY)

    markdown = (tmp_path / 'Q5.md').read_text(encoding='utf-8')
    docx_text = '\n'.join(paragraph.text for paragraph in Document(tmp_path / 'Q5.docx').paragraphs)
    assert missing_prefixes(markdown) == []
    assert missing_prefixes(docx_text) == []
