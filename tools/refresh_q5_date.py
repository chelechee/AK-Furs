"""Обновление обеих дат Q5 (записка руководителю и ответ клиентке) на текущий день."""

import sys
from datetime import date
from pathlib import Path
import re

from docx import Document

ROOT = Path(__file__).resolve().parents[1]

PREFIXES = (
    'AK Furs · итоги сверки данных · ',
    'AK Furs · ответ менеджера · ',
)
DATE_PATTERN = re.compile(r'\d{2}\.\d{2}\.\d{4}')


def replace_dates(text, today):
    updated = text
    for prefix in PREFIXES:
        updated = re.sub(
            '(' + re.escape(prefix) + r')\d{2}\.\d{2}\.\d{4}',
            lambda match: match.group(1) + today,
            updated,
        )
    return updated


def missing_prefixes(text):
    return [prefix for prefix in PREFIXES
            if not re.search(re.escape(prefix) + r'\d{2}\.\d{2}\.\d{4}', text)]


def update_markdown(path, today):
    path = Path(path)
    if not path.exists():
        return False
    text = path.read_text(encoding='utf-8')
    updated = replace_dates(text, today)
    if updated != text:
        path.write_text(updated, encoding='utf-8', newline='\n')
        return True
    return False


def _set_prefix_text(paragraph, prefix, today):
    if paragraph.runs:
        paragraph.runs[0].text = prefix + today
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(prefix + today)


def update_docx(path, today):
    path = Path(path)
    if not path.exists():
        return False
    document = Document(path)
    changed = False
    for paragraph in document.paragraphs:
        for prefix in PREFIXES:
            if paragraph.text.startswith(prefix) and paragraph.text != prefix + today:
                _set_prefix_text(paragraph, prefix, today)
                changed = True
    if changed:
        document.save(path)
    return changed


def targets():
    return [
        ROOT / 'Q5.md',
        ROOT / 'Q5.docx',
        ROOT / 'employer_package' / 'Q5.docx',
    ]


def main():
    today = date.today().strftime('%d.%m.%Y')
    for path in targets():
        if not path.exists():
            print('{}: not found'.format(path.relative_to(ROOT)))
            continue
        if path.suffix == '.md':
            changed = update_markdown(path, today)
            missing = missing_prefixes(path.read_text(encoding='utf-8'))
        else:
            changed = update_docx(path, today)
            missing = missing_prefixes('\n'.join(p.text for p in Document(path).paragraphs))
        print('{}: {}{}'.format(
            path.relative_to(ROOT),
            'updated' if changed else 'unchanged',
            '' if not missing else ' (missing dates: {})'.format(', '.join(missing))))
    print('date:', today)
    return 0


if __name__ == '__main__':
    sys.exit(main())
