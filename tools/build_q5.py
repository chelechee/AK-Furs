"""Генерация Q5 из данных: записка руководителю и ответ клиентке."""

import sys
from datetime import date, datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from akfurs import q5
from akfurs.io_loaders import load_dataset
from solve import _data_source

DEFAULT_SOURCE = Path('data')
DEFAULT_MARKDOWN = ROOT / 'Q5.md'
DEFAULT_DOCX = ROOT / 'Q5.docx'
ANCHORS = ((0, 'Записка'), (2, 'Тема:'), (8, 'Что предлагаю'), (12, 'Ответ клиентке'))
FONT_NAME = 'Arial'
BODY_COLOR = RGBColor(48, 47, 45)
HEADING_COLOR = RGBColor(47, 43, 38)
ACCENT_COLOR = RGBColor(125, 102, 70)
MUTED_COLOR = RGBColor(112, 109, 104)


def _set_run_font(run, size, *, bold=False, italic=False, color=BODY_COLOR):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        properties.insert(0, fonts)
    for attribute in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + attribute), FONT_NAME)


def format_q5_document(document, texts):
    normal = document.styles['Normal']
    normal.font.name = FONT_NAME
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BODY_COLOR
    normal_properties = normal._element.get_or_add_rPr()
    normal_fonts = normal_properties.rFonts
    if normal_fonts is None:
        normal_fonts = OxmlElement('w:rFonts')
        normal_properties.insert(0, normal_fonts)
    for attribute in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        normal_fonts.set(qn('w:' + attribute), FONT_NAME)

    for section in document.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.header_distance = Cm(0.8)
        section.footer_distance = Cm(0.8)

    for index, (paragraph, text) in enumerate(zip(document.paragraphs, texts)):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.line_spacing = 1.08
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.left_indent = Cm(0)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.keep_with_next = index in (0, 1, 2, 8, 12, 13)
        paragraph.paragraph_format.page_break_before = index == 12

        size, bold, italic, color = 10.5, False, False, BODY_COLOR
        if index in (0, 12):
            size, bold, color = (15 if index == 0 else 14), True, HEADING_COLOR
            paragraph.paragraph_format.space_before = Pt(0 if index == 0 else 0)
            paragraph.paragraph_format.space_after = Pt(7)
        elif index in (1, 13):
            size, italic, color = 9, True, MUTED_COLOR
            paragraph.paragraph_format.space_after = Pt(11)
        elif index == 2:
            size, bold, color = 11, True, ACCENT_COLOR
            paragraph.paragraph_format.space_after = Pt(8)
        elif index == 8:
            size, bold, color = 11, True, ACCENT_COLOR
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(4)
        elif index in (5, 6, 7, 9, 10):
            paragraph.paragraph_format.left_indent = Cm(0.45)
            paragraph.paragraph_format.first_line_indent = Cm(-0.45)
            paragraph.paragraph_format.space_after = Pt(5)
        elif index == 11:
            paragraph.paragraph_format.space_after = Pt(0)
        elif index >= 14:
            size = 11
            paragraph.paragraph_format.space_after = Pt(9)

        if not paragraph.runs and text:
            paragraph.add_run(text)
        for run in paragraph.runs:
            _set_run_font(run, size, bold=bold, italic=italic, color=color)

    document.core_properties.title = 'Записка руководителю и ответ клиентке — AK Furs'
    document.core_properties.subject = 'Сверка данных и ответ клиентке C0010'
    document.core_properties.author = 'AK Furs'


def update_docx(path, texts):
    document = Document(path)
    paragraphs = document.paragraphs
    missing_count = len(texts) - len(paragraphs)
    if missing_count > 0:
        if (len(paragraphs) < 15 or not paragraphs[14].text.startswith(
                ('Ответ не сформирован:', 'Здравствуйте!')) or missing_count > 3):
            raise ValueError('{} has {} paragraphs, expected {}'.format(
                path.name, len(paragraphs), len(texts)))
        for paragraph in paragraphs[14:]:
            element = paragraph._element
            element.getparent().remove(element)
        for text in texts[14:]:
            document.add_paragraph(text)
        paragraphs = document.paragraphs

    extra_count = len(paragraphs) - len(texts)
    if extra_count > 2:
        raise ValueError('{} has {} paragraphs, expected {}'.format(
            path.name, len(paragraphs), len(texts)))
    for index, prefix in ANCHORS:
        if not paragraphs[index].text.startswith(prefix):
            raise ValueError('{}: paragraph {} does not start with {!r}'.format(
                path.name, index, prefix))
    legacy_extras = [
        'Я дополнительно проверю актуальность цены, скидки и наличие до вашего визита.',
        'Справочно для проверки перед отправкой:',
    ]
    if extra_count:
        actual_extras = [paragraph.text for paragraph in paragraphs[-extra_count:]]
        old_client_text = actual_extras[-1].startswith((
            legacy_extras[-1], 'Уточню актуальную стоимость и наличие', 'В заказе указана цена',
        ))
        if old_client_text and (
                extra_count == 1 or actual_extras[0].startswith(legacy_extras[0])):
            for paragraph in paragraphs[-extra_count:]:
                element = paragraph._element
                element.getparent().remove(element)
        else:
            raise ValueError('{} has unexpected extra paragraphs'.format(path.name))
        paragraphs = document.paragraphs
    for paragraph, text in zip(paragraphs, texts):
        if paragraph.runs:
            paragraph.runs[0].text = text
            for run in paragraph.runs[1:]:
                run.text = ''
        elif text:
            paragraph.add_run(text)
    format_q5_document(document, texts)
    document.save(path)


def render_q5(dataset, today=None):
    """Prepare both Q5 formats from an already-loaded dataset."""
    today = today or date.today().strftime('%d.%m.%Y')
    figures = q5.build_figures(dataset)
    text = q5.markdown(figures, today)
    stale = q5.unexpected_numbers(text, q5.allowed_numbers(figures))
    if stale:
        raise ValueError('Q5 contains numbers that are not derived from data: {}'.format(
            ', '.join(stale)))
    return figures, text, q5.docx_texts(figures, today)


def write_q5(markdown_path, docx_path, markdown_text, docx_texts):
    markdown_path = Path(markdown_path)
    markdown_path.parent.mkdir(parents=True, exist_ok=True)
    markdown_path.write_text(markdown_text, encoding='utf-8', newline='\n')
    docx_path = Path(docx_path)
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    if docx_path.exists():
        update_docx(docx_path, docx_texts)
    else:
        document = Document()
        for text in docx_texts:
            document.add_paragraph(text)
        format_q5_document(document, docx_texts)
        document.save(docx_path)


def parse_args(argv):
    args = list(argv[1:])
    options = {'source': DEFAULT_SOURCE, 'today': date.today().strftime('%d.%m.%Y'),
               'markdown': DEFAULT_MARKDOWN, 'docx': DEFAULT_DOCX}
    flags = {'--date': 'today', '--out': 'markdown', '--docx': 'docx'}
    while args:
        head = args.pop(0)
        if head in flags:
            if not args:
                raise SystemExit('{} requires a value'.format(head))
            value = args.pop(0)
            if head == '--date':
                datetime.strptime(value, '%d.%m.%Y')
            options[flags[head]] = Path(value) if head != '--date' else value
        elif head.startswith('--'):
            raise SystemExit('unknown option: {}'.format(head))
        else:
            options['source'] = Path(head)
    return options


def main(argv):
    try:
        options = parse_args(argv)
        with _data_source(options['source']) as data_dir:
            dataset = load_dataset(data_dir)
        figures, text, paragraphs = render_q5(dataset, options['today'])
        write_q5(options['markdown'], options['docx'], text, paragraphs)
    except Exception as error:
        print('ERROR: {}: {}'.format(type(error).__name__, error), file=sys.stderr)
        return 1

    print('OK: {} and {} generated for {}'.format(
        options['markdown'].name, options['docx'].name, options['today']))
    print('Q1 {}/{}; Q2a {}; Q2b {}/{}; Q3 {}; Q4 {}/{}'.format(
        figures['unique_orders'], figures['unique_lines'],
        figures['paid_without_payment'], figures['unidentified_count'],
        figures['unidentified_sum'].replace('₽', 'RUB'), figures['violations_total'],
        figures['flagged_chats'], figures['chats_total']))
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv))
