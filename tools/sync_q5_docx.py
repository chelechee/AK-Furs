"""Synchronize the C0010 reply in Q5 Word deliverables from Q5.md."""

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
HEADING = 'ОТВЕТ КЛИЕНТКЕ В ЧАТЕ C0010'


def _reply_from_markdown(path):
    text = path.read_text(encoding='utf-8')
    if HEADING not in text:
        raise ValueError('Q5 client-reply heading not found in {}'.format(path))
    section = text.split(HEADING, 1)[1]
    section = section.split('Справочно для проверки перед отправкой:', 1)[0]
    reply = [line.strip() for line in section.splitlines() if line.strip()]
    if not reply:
        raise ValueError('Q5 client reply is empty in {}'.format(path))
    return reply


def _set_paragraph_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ''
    elif text:
        paragraph.add_run(text)


def sync_docx(markdown_path, docx_path):
    reply = _reply_from_markdown(markdown_path)
    document = Document(docx_path)
    start = next((index for index, paragraph in enumerate(document.paragraphs)
                  if paragraph.text.startswith('Здравствуйте!')), None)
    if start is None:
        raise ValueError('client reply paragraph not found in {}'.format(docx_path))

    slots = []
    for paragraph in document.paragraphs[start:]:
        if paragraph.text.startswith('Справочно для проверки перед отправкой:'):
            break
        if paragraph.text.strip():
            slots.append(paragraph)
    if len(slots) < len(reply):
        raise ValueError('not enough existing reply paragraphs in {}'.format(docx_path))

    for index, paragraph in enumerate(slots):
        _set_paragraph_text(paragraph, reply[index] if index < len(reply) else '')
    document.save(docx_path)


def main():
    targets = (
        (ROOT / 'Q5.md', ROOT / 'Q5.docx'),
        (ROOT / 'handoff' / 'Q5.md', ROOT / 'handoff' / 'Q5.docx'),
    )
    for markdown_path, docx_path in targets:
        if markdown_path.exists() and docx_path.exists():
            sync_docx(markdown_path, docx_path)
            print('updated:', docx_path.relative_to(ROOT))


if __name__ == '__main__':
    main()
