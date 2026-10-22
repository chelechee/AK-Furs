"""Генератор стресс-вариантов набора данных AK Furs.

Три группы:
  INVARIANT — данные переписаны в другой форме, ответы обязаны совпасть с базовыми;
  MUTATIONS — данные изменены осмысленно, ответы обязаны измениться предсказуемо;
  BROKEN    — данные повреждены, конвейер обязан упасть с понятной ошибкой.
"""

import csv
import io
import random
import re
import shutil
import sys
import zipfile

import openpyxl
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from akfurs.normalize import parse_date  # noqa: E402

CRM = '01_crm_orders_export.csv'
BANK = '02_bank_statement_rub.csv'
PRICEBOOK = '03_catalog_prices_reference.xlsx'
CHATS = '04_client_chats.zip'

PRICE_COLUMN = 8
DATE_COLUMN = 2
ORDER_COLUMN = 0
STATUS_COLUMN = 11


def _prepare(src, dest):
    dest = Path(dest)
    dest.mkdir(parents=True, exist_ok=True)
    for name in (CRM, BANK, PRICEBOOK, CHATS):
        shutil.copyfile(Path(src) / name, dest / name)
    return dest


def _rows(path, encoding, delimiter=';'):
    text = Path(path).read_bytes().decode(encoding)
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter) if row]


def _crm(src):
    return _rows(Path(src) / CRM, 'utf-8-sig')


def _bank(src):
    return _rows(Path(src) / BANK, 'cp1251')


def _write(path, rows, encoding='utf-8', delimiter=';', lineterminator='\n'):
    with open(path, 'w', encoding=encoding, newline='') as handle:
        csv.writer(handle, delimiter=delimiter, lineterminator=lineterminator).writerows(rows)


def _rewrite_chats(src, dest, encoding):
    with zipfile.ZipFile(Path(src) / CHATS) as reader, zipfile.ZipFile(Path(dest) / CHATS, 'w') as writer:
        for info in reader.infolist():
            text = reader.read(info).decode('utf-8')
            writer.writestr(info.filename, text.encode(encoding))


def _money_to_int(text):
    value = text.strip().replace('\xa0', ' ').replace('\u202f', ' ').replace('\u2009', ' ')
    if 'млн' in value:
        raise ValueError('skip millions')
    value = value.replace(' ', '')
    if ',' in value and '.' in value:
        value = value.replace(',', '')
    elif ',' in value:
        if re.fullmatch(r'\d{1,3}(,\d{3})+', value):
            value = value.replace(',', '')
        else:
            value = value.replace(',', '.')
    return Decimal(value)


def _european(value):
    text = f'{value:,.2f}'
    return text.replace(',', '\x00').replace('.', ',').replace('\x00', '.')


def variant_encodings(src, dest):
    dest = _prepare(src, dest)
    rows = [[cell.replace('₽', 'руб') for cell in row] for row in _crm(src)]
    _write(dest / CRM, rows, encoding='cp1251')
    _write(dest / BANK, _bank(src), encoding='utf-8')
    _rewrite_chats(src, dest, 'utf-8-sig')


def variant_delimiters(src, dest):
    dest = _prepare(src, dest)
    _write(dest / CRM, _crm(src), delimiter=',')
    _write(dest / BANK, _bank(src), delimiter='\t', encoding='cp1251')


def variant_crlf(src, dest):
    dest = _prepare(src, dest)
    _write(dest / CRM, _crm(src), lineterminator='\r\n')
    _write(dest / BANK, _bank(src), lineterminator='\r\n', encoding='cp1251')


def variant_columns(src, dest):
    dest = _prepare(src, dest)
    _write(dest / CRM, [list(reversed(row)) for row in _crm(src)])
    _write(dest / BANK, [list(reversed(row)) for row in _bank(src)], encoding='cp1251')


def variant_dates_alt(src, dest):
    dest = _prepare(src, dest)
    rows = _crm(src)
    patterns = ('%d.%m.%Y', '%d-%m-%Y', '%Y/%m/%d', '%Y-%m-%d 00:00:00')
    for index, row in enumerate(rows[1:]):
        row[DATE_COLUMN] = parse_date(row[DATE_COLUMN]).strftime(patterns[index % len(patterns)])
    _write(dest / CRM, rows)


def variant_prices_eu(src, dest):
    dest = _prepare(src, dest)
    rows = _crm(src)
    for row in rows[1:]:
        try:
            row[PRICE_COLUMN] = _european(_money_to_int(row[PRICE_COLUMN]))
        except (ValueError, ArithmeticError):
            continue
    _write(dest / CRM, rows)

    bank = _bank(src)
    for row in bank[2:]:
        row[2] = _european(_money_to_int(row[2]))
    _write(dest / BANK, bank, encoding='cp1251')


def variant_spaces(src, dest):
    dest = _prepare(src, dest)
    rows = _crm(src)
    for index, row in enumerate(rows[1:]):
        separator = '\u00a0' if index % 2 == 0 else '\u202f'
        row[PRICE_COLUMN] = row[PRICE_COLUMN].replace(' ', separator)
    _write(dest / CRM, rows)

    bank = _bank(src)
    for row in bank[2:]:
        row[2] = row[2].replace(' ', '\u00a0')
    _write(dest / BANK, bank, encoding='cp1251')


def variant_numbers(src, dest):
    dest = _prepare(src, dest)
    rows = _crm(src)
    patterns = ('№{}', 'заказ {}', 'A\u2013{}', ' A {}', 'A{}')
    for index, row in enumerate(rows[1:]):
        digits = re.sub(r'\D', '', row[ORDER_COLUMN])
        row[ORDER_COLUMN] = patterns[index % len(patterns)].format(digits)
    _write(dest / CRM, rows)


def variant_shuffle(src, dest):
    dest = _prepare(src, dest)
    rows = _crm(src)
    body = rows[1:]
    random.Random(20261022).shuffle(body)
    _write(dest / CRM, [rows[0]] + body)


def variant_late_duplicate(src, dest):
    dest = _prepare(src, dest)
    rows = _crm(src)
    duplicate = list(rows[1])
    duplicate[len(duplicate) - 1] = '2026-10-02 00:00'
    _write(dest / CRM, rows + [duplicate])


def mutation_cover_unpaid_order(src, dest):
    dest = _prepare(src, dest)
    shutil.copyfile(Path(src) / CRM, dest / CRM)
    bank = _bank(src)
    bank.append(['01.10.2026', '12:00', '1 000,00', 'RUR', 'Поступление',
                 'Тестовая оплата', 'Оплата заказа A-00012'])
    _write(dest / BANK, bank, encoding='cp1251')


def mutation_add_unidentified_payment(src, dest):
    dest = _prepare(src, dest)
    shutil.copyfile(Path(src) / CRM, dest / CRM)
    bank = _bank(src)
    bank.append(['01.10.2026', '12:00', '100 000,00', 'RUR', 'Поступление',
                 'Тестовый перевод', 'Оплата по договору'])
    _write(dest / BANK, bank, encoding='cp1251')


def mutation_add_cheap_line(src, dest):
    dest = _prepare(src, dest)
    shutil.copyfile(Path(src) / BANK, dest / BANK)
    rows = _crm(src)
    rows.append(['A-00001', '9', '2025-05-05', 'CL-00001', 'Вероника Белова', 'Сайт',
                 'CHN-PAL-01', '1', '1000', 'RUB', '', 'Оплачен', '', '2026-10-02 00:00'])
    _write(dest / CRM, rows)


def broken_ambiguous_payment(src, dest):
    dest = _prepare(src, dest)
    shutil.copyfile(Path(src) / CRM, dest / CRM)
    bank = _bank(src)
    bank.append(['01.10.2026', '12:00', '1 000,00', 'RUR', 'Поступление',
                 'Спорный платёж', 'заказ 12 и 34'])
    _write(dest / BANK, bank, encoding='cp1251')


def broken_unknown_date(src, dest):
    dest = _prepare(src, dest)
    shutil.copyfile(Path(src) / BANK, dest / BANK)
    rows = _crm(src)
    rows[1][DATE_COLUMN] = 'вчера'
    _write(dest / CRM, rows)


def variant_renamed_sheets(src, dest):
    dest = _prepare(src, dest)
    workbook = openpyxl.load_workbook(Path(src) / PRICEBOOK)
    names = {'Каталог': 'Catalog', 'Переименования': 'Renames',
             'Прайс_история': 'PriceHistory', 'Менеджеры': 'Managers'}
    for old, new in names.items():
        workbook[old].title = new
    workbook['PriceHistory'].sheet_state = 'veryHidden'
    workbook._sheets = list(reversed(workbook._sheets))
    workbook.save(dest / PRICEBOOK)


def broken_formula_prices(src, dest):
    dest = _prepare(src, dest)
    workbook = openpyxl.load_workbook(Path(src) / PRICEBOOK)
    sheet = workbook['Прайс_история']
    for row in sheet.iter_rows(min_row=2, min_col=4, max_col=5):
        for cell in row:
            cell.value = '={}*1'.format(cell.value)
    workbook.save(dest / PRICEBOOK)


INVARIANT_VARIANTS = {
    'encodings': variant_encodings,
    'delimiters': variant_delimiters,
    'crlf': variant_crlf,
    'columns': variant_columns,
    'dates_alt': variant_dates_alt,
    'prices_eu': variant_prices_eu,
    'spaces': variant_spaces,
    'numbers': variant_numbers,
    'shuffle': variant_shuffle,
    'late_duplicate': variant_late_duplicate,
    'renamed_sheets': variant_renamed_sheets,
}

MUTATIONS = {
    'cover_unpaid_order': mutation_cover_unpaid_order,
    'add_unidentified_payment': mutation_add_unidentified_payment,
    'add_cheap_line': mutation_add_cheap_line,
}

BROKEN = {
    'ambiguous_payment': broken_ambiguous_payment,
    'unknown_date': broken_unknown_date,
    'formula_prices': broken_formula_prices,
}


def build_all(src, out):
    out = Path(out)
    report = []
    for name, builder in {**INVARIANT_VARIANTS, **MUTATIONS, **BROKEN}.items():
        destination = out / name
        builder(src, destination)
        report.append('{} -> {}'.format(name, destination))
    return report


if __name__ == '__main__':
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'data'
    target = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / 'stress'
    for line in build_all(source, target):
        print(line)
